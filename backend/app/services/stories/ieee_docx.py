"""Build IEEE-style Word (.docx) exports for Write stories."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from ...models.stories import IEEE_SECTIONS, Story


def _set_run_font(run, *, name: str = "Times New Roman", size_pt: float = 10, bold: bool = False, italic: bool = False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size_pt)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def _add_centered(doc: Document, text: str, *, size_pt: float, bold: bool = False, italic: bool = False, space_after: float = 6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    _set_run_font(run, size_pt=size_pt, bold=bold, italic=italic)
    return p


def _add_heading_ieee(doc: Document, text: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    _set_run_font(run, size_pt=10, bold=True)


def _add_body_paragraphs(doc: Document, text: str, *, first_line_indent: bool = True):
    chunks = re.split(r"\n\s*\n", (text or "").strip())
    if not chunks:
        return
    for i, chunk in enumerate(chunks):
        lines = [ln.strip() for ln in chunk.splitlines() if ln.strip()]
        if not lines:
            continue
        # Reference-style lines: keep each on its own paragraph without indent.
        if all(re.match(r"^\[\d+\]", ln) for ln in lines):
            for ln in lines:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.first_line_indent = Pt(0)
                run = p.add_run(ln)
                _set_run_font(run, size_pt=9)
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0
        if first_line_indent and i > 0:
            p.paragraph_format.first_line_indent = Inches(0.2)
        run = p.add_run(" ".join(lines))
        _set_run_font(run, size_pt=10)


def build_ieee_docx(story: Story) -> bytes:
    """Return .docx bytes for an IEEE-style research paper draft."""

    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    # IEEE conference-ish margins (approx. 0.75").
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    title = (story.title or "Untitled").strip()
    _add_centered(doc, title, size_pt=24, bold=True, space_after=12)

    authors = (story.authors_line or story.author_name or "").strip()
    if authors:
        _add_centered(doc, authors, size_pt=11, space_after=4)

    affiliation = (story.affiliation or "").strip()
    if affiliation:
        _add_centered(doc, affiliation, size_pt=10, italic=True, space_after=14)

    sections = story.sections or {}
    if story.format == "ieee" or any((sections.get(k) or "").strip() for k, _ in IEEE_SECTIONS):
        for key, label in IEEE_SECTIONS:
            text = (sections.get(key) or "").strip()
            if not text and key not in ("abstract", "introduction"):
                # Still show empty core sections so the template is complete when exporting drafts? 
                # Skip empty to keep download clean.
                continue
            if not text:
                continue
            if key == "abstract":
                _add_heading_ieee(doc, "Abstract")
                _add_body_paragraphs(doc, text, first_line_indent=False)
            elif key == "keywords":
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(10)
                run_l = p.add_run("Index Terms—")
                _set_run_font(run_l, size_pt=10, italic=True, bold=True)
                run_v = p.add_run(text)
                _set_run_font(run_v, size_pt=10, italic=True)
            else:
                _add_heading_ieee(doc, label)
                _add_body_paragraphs(doc, text, first_line_indent=key != "references")
    else:
        # Freeform fallback: dump body_md as paragraphs under the title block.
        body = (story.body_md or "").strip()
        if body:
            _add_body_paragraphs(doc, body, first_line_indent=False)

    footer = doc.sections[0].footer
    fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Prepared in Lexicon Gate · IEEE-style draft")
    _set_run_font(fr, size_pt=8, italic=True)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def safe_docx_filename(title: str) -> str:
    base = re.sub(r"[^A-Za-z0-9._-]+", "_", (title or "paper").strip())[:60].strip("._") or "paper"
    if not base.lower().endswith(".docx"):
        base = f"{base}.docx"
    return base
