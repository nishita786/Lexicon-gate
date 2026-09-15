"""Export IEEE-style paper drafts to DOCX and PDF."""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from ...models.paper_drafts import SECTION_LABELS, SECTION_ORDER, PaperDraft
from .ieee_html import render_ieee_html


def build_preview_html(draft: PaperDraft) -> str:
    return render_ieee_html(draft, mode="preview")


def build_docx(draft: PaperDraft) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.orientation = WD_ORIENT.PORTRAIT
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, margin, Inches(0.75))

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(draft.title or "Untitled draft")
    run.bold = True
    run.font.size = Pt(18)
    run.font.name = "Times New Roman"

    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ar = authors.add_run(draft.authors or "Author")
    ar.font.size = Pt(11)
    ar.font.name = "Times New Roman"

    sections = draft.sections or {}
    for key in ("abstract", "keywords"):
        _add_heading(doc, SECTION_LABELS.get(key, key.title()), italic=True)
        for para in _split_paras(sections.get(key) or ""):
            p = doc.add_paragraph(para)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9)

    # Continuous section break, then two-column body (IEEE-style).
    new_section = doc.add_section()
    new_section.page_width = Inches(8.5)
    new_section.page_height = Inches(11)
    for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(new_section, margin, Inches(0.75))
    _set_section_start_continuous(new_section)
    _set_section_columns(new_section, num=2, space_twips=360)

    for key in SECTION_ORDER:
        if key in ("abstract", "keywords"):
            continue
        _add_heading(doc, SECTION_LABELS.get(key, key.title()))
        for para in _split_paras(sections.get(key) or ""):
            p = doc.add_paragraph(para)
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)

    if draft.references:
        _add_heading(doc, "References")
        for ref in draft.references:
            p = doc.add_paragraph(ref)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_pdf(draft: PaperDraft) -> bytes:
    from xhtml2pdf import pisa

    html = render_ieee_html(draft, mode="pdf")
    buffer = io.BytesIO()
    result = pisa.CreatePDF(html, dest=buffer, encoding="utf-8")
    if result.err:
        raise RuntimeError("PDF export failed.")
    return buffer.getvalue()


def _add_heading(doc: Document, text: str, *, italic: bool = False) -> None:
    heading = doc.add_paragraph()
    hr = heading.add_run(text)
    hr.bold = True
    hr.italic = italic
    hr.font.size = Pt(10)
    hr.font.name = "Times New Roman"
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(2)


def _set_section_start_continuous(section) -> None:
    sectPr = section._sectPr  # noqa: SLF001
    typ = sectPr.find(qn("w:type"))
    if typ is None:
        typ = OxmlElement("w:type")
        sectPr.insert(0, typ)
    typ.set(qn("w:val"), "continuous")


def _set_section_columns(section, *, num: int, space_twips: int) -> None:
    sectPr = section._sectPr  # noqa: SLF001
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space_twips))
    cols.set(qn("w:equalWidth"), "1")


def _split_paras(text: str) -> list[str]:
    chunks = re.split(r"\n\s*\n", (text or "").strip())
    out = [c.strip() for c in chunks if c.strip()]
    return out or [""]
